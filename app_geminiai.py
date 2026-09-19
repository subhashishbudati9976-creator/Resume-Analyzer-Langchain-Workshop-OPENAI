"""
Resume Analyzer & Career Toolkit
Run:
    pip install streamlit PyPDF2 langchain-openai langchain-core reportlab python-docx google-genai
    streamlit run app.py
"""

import io
import json
import re
import sqlite3
from datetime import datetime
from typing import Any, Optional

import streamlit as st
from PyPDF2 import PdfReader
from langchain_openai import ChatOpenAI
from langchain_core.messages import SystemMessage, HumanMessage

try:
    from google import genai
    GEMINI_AVAILABLE = True
except Exception:
    genai = None
    GEMINI_AVAILABLE = False


# Optional document exporters
try:
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    REPORTLAB_AVAILABLE = True
except Exception:
    REPORTLAB_AVAILABLE = False

try:
    from docx import Document
    from docx.shared import Pt
    DOCX_AVAILABLE = True
except Exception:
    DOCX_AVAILABLE = False


st.set_page_config(page_title="Resume Analyzer & Career Toolkit", page_icon="📄", layout="wide")

DB_PATH = "resume_toolkit.db"


# ----------------------------- Storage ---------------------------------------
def init_db():
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS saved_resumes (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                resume_text TEXT NOT NULL,
                created_at TEXT NOT NULL
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS applications (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                company TEXT NOT NULL,
                role TEXT NOT NULL,
                status TEXT NOT NULL,
                date_applied TEXT,
                notes TEXT
            )
        """)


def save_resume(name: str, text: str):
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute(
            "INSERT INTO saved_resumes(name, resume_text, created_at) VALUES (?, ?, ?)",
            (name, text, datetime.now().strftime("%Y-%m-%d %H:%M")),
        )


def get_saved_resumes():
    with sqlite3.connect(DB_PATH) as conn:
        return conn.execute(
            "SELECT id, name, created_at FROM saved_resumes ORDER BY id DESC"
        ).fetchall()


def get_resume_text(resume_id: int):
    with sqlite3.connect(DB_PATH) as conn:
        row = conn.execute(
            "SELECT resume_text FROM saved_resumes WHERE id=?", (resume_id,)
        ).fetchone()
    return row[0] if row else ""


def delete_resume(resume_id: int):
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute("DELETE FROM saved_resumes WHERE id=?", (resume_id,))


def add_application(company, role, status, date_applied, notes):
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute(
            "INSERT INTO applications(company, role, status, date_applied, notes) VALUES (?, ?, ?, ?, ?)",
            (company, role, status, date_applied, notes),
        )


def get_applications():
    with sqlite3.connect(DB_PATH) as conn:
        return conn.execute(
            "SELECT id, company, role, status, date_applied, notes FROM applications ORDER BY id DESC"
        ).fetchall()


def update_application(app_id, status, notes):
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute(
            "UPDATE applications SET status=?, notes=? WHERE id=?",
            (status, notes, app_id),
        )


def delete_application(app_id):
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute("DELETE FROM applications WHERE id=?", (app_id,))


init_db()


# ----------------------------- Helpers ---------------------------------------
def extract_pdf_text(uploaded_file) -> str:
    try:
        reader = PdfReader(uploaded_file)
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        return "\n".join(pages).strip()
    except Exception as exc:
        raise ValueError(f"Could not read this PDF: {exc}") from exc


def clean_json_text(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s*```$", "", text)
    return text.strip()


def parse_json_response(text: str) -> dict:
    cleaned = clean_json_text(text)
    try:
        result = json.loads(cleaned)
        if isinstance(result, dict):
            return result
    except json.JSONDecodeError:
        pass

    # Recover a JSON object if the model wrapped it in additional prose.
    start, end = cleaned.find("{"), cleaned.rfind("}")
    if start >= 0 and end > start:
        result = json.loads(cleaned[start:end + 1])
        if isinstance(result, dict):
            return result
    raise ValueError("The AI response was not valid JSON. Please try again.")


def call_ai(system_prompt: str, user_prompt: str, output_parser=None, *args, **kwargs):
    """
    Route the request to the selected provider using direct messages.
    Accepts an optional third positional argument for compatibility.
    """
    provider = st.session_state.get("ai_provider", "OpenAI")

    if provider == "Google Gemini":
        if not GEMINI_AVAILABLE:
            raise RuntimeError(
                "Gemini support requires the google-genai package. "
                "Install it with: pip install google-genai"
            )
        api_key = st.session_state.get("gemini_api_key", "").strip()
        model_name = st.session_state.get("gemini_model_name", "gemini-2.5-flash").strip()
        if not api_key:
            raise ValueError("Add your Google Gemini API key in the sidebar.")

        client = genai.Client(api_key=api_key)
        response = client.models.generate_content(
            model=model_name,
            contents=[
                f"SYSTEM INSTRUCTIONS:\n{system_prompt}\n\nUSER REQUEST:\n{user_prompt}"
            ],
        )
        content = response.text or ""
    else:
        api_key = st.session_state.get("api_key", "").strip()
        model_name = st.session_state.get("model_name", "gpt-4o-mini").strip()
        if not api_key:
            raise ValueError("Add your OpenAI API key in the sidebar.")

        llm = ChatOpenAI(
            model=model_name,
            api_key=api_key,
            temperature=0.2,
        )
        response = llm.invoke([
            SystemMessage(content=str(system_prompt)),
            HumanMessage(content=str(user_prompt)),
        ])
        content = response.content
        if isinstance(content, list):
            content = "\n".join(
                part.get("text", "") if isinstance(part, dict) else str(part)
                for part in content
            )
        content = str(content)

    if output_parser is not None and hasattr(output_parser, "parse"):
        return output_parser.parse(content)
    return content



def safe_list(value):
    return value if isinstance(value, list) else []


def markdown_to_pdf(text: str) -> bytes:
    if not REPORTLAB_AVAILABLE:
        raise RuntimeError("Install reportlab to export PDF: pip install reportlab")
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4, rightMargin=42, leftMargin=42, topMargin=42, bottomMargin=42
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ResumeTitle", parent=styles["Title"], fontName="Helvetica-Bold",
        fontSize=17, leading=21, textColor=colors.HexColor("#202020"), spaceAfter=12
    )
    heading_style = ParagraphStyle(
        "ResumeHeading", parent=styles["Heading2"], fontSize=11,
        leading=14, spaceBefore=10, spaceAfter=4, textColor=colors.HexColor("#333333")
    )
    body_style = ParagraphStyle(
        "ResumeBody", parent=styles["BodyText"], fontSize=9,
        leading=12, spaceAfter=3
    )

    story = []
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            story.append(Spacer(1, 4))
            continue
        escaped = (line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))
        if line.startswith("# "):
            story.append(Paragraph(escaped[2:], title_style))
        elif line.startswith("## ") or line.startswith("### "):
            story.append(Paragraph(escaped.lstrip("# ").strip(), heading_style))
        else:
            escaped = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", escaped)
            if escaped.startswith("- ") or escaped.startswith("* "):
                escaped = "• " + escaped[2:]
            story.append(Paragraph(escaped, body_style))
    doc.build(story)
    return buffer.getvalue()


def markdown_to_docx(text: str) -> bytes:
    if not DOCX_AVAILABLE:
        raise RuntimeError("Install python-docx to export DOCX: pip install python-docx")
    doc = Document()
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:], level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith(("- ", "* ")):
            doc.add_paragraph(line[2:], style="List Bullet")
        else:
            p = doc.add_paragraph()
            p.add_run(re.sub(r"\*\*(.*?)\*\*", r"\1", line))
            for run in p.runs:
                run.font.size = Pt(10)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def display_list(items, empty_message="None identified."):
    items = safe_list(items)
    if not items:
        st.write(empty_message)
    else:
        for item in items:
            if isinstance(item, dict):
                title = item.get("skill") or item.get("name") or item.get("title") or "Item"
                detail = item.get("reason") or item.get("evidence") or item.get("details") or ""
                st.markdown(f"- **{title}**" + (f" — {detail}" if detail else ""))
            else:
                st.markdown(f"- {item}")


# ----------------------------- Sidebar ---------------------------------------
st.sidebar.title("Resume Toolkit")
st.sidebar.caption("Analyze, tailor, prepare, and track.")
st.sidebar.selectbox(
    "AI provider",
    ["OpenAI", "Google Gemini"],
    key="ai_provider",
)
st.sidebar.caption("Choose which provider handles AI requests.")

if st.session_state.get("ai_provider", "OpenAI") == "Google Gemini":
    st.sidebar.text_input("Google Gemini API key", type="password", key="gemini_api_key")
    st.sidebar.text_input(
        "Gemini model",
        value="gemini-2.5-flash",
        key="gemini_model_name",
    )
    st.sidebar.caption("Get a Gemini API key from Google AI Studio. Install google-genai if needed.")
else:
    st.sidebar.text_input("OpenAI API key", type="password", key="api_key")
    st.sidebar.text_input("OpenAI model", value="gpt-4o-mini", key="model_name")
    st.sidebar.caption("Use a model name enabled for your OpenAI API account.")

page = st.sidebar.radio(
    "Navigate",
    ["Resume Analyzer & Builder", "Interview Preparation", "Skill Roadmap", "Application Tracker", "Saved Resumes"],
)

st.title("📄 Resume Analyzer & Career Toolkit")
st.caption("Use AI to review your resume and tailor it to a role. Always verify every generated claim before using it.")


# ----------------------------- Main Analyzer ---------------------------------
if page == "Resume Analyzer & Builder":
    left, right = st.columns([1, 1])
    with left:
        uploaded_pdf = st.file_uploader("Upload your current resume (PDF)", type=["pdf"])
        job_description = st.text_area(
            "Paste the job description",
            height=230,
            placeholder="Paste the full job description here...",
        )
        candidate_name = st.text_input("Name to use on tailored resume (optional)")
        verified_info = st.text_area(
            "Additional verified details (optional)",
            height=120,
            placeholder="Add only accurate projects, skills, education, achievements, links, or experience not present in the PDF.",
        )
        fresher_mode = st.checkbox("I'm applying as a student / fresher", value=True)
        template = st.selectbox(
            "Resume style",
            ["Classic ATS", "Modern Professional", "Student / Fresher", "Minimal"],
        )
        analyze_clicked = st.button("Analyze Resume", type="primary", use_container_width=True)

    if uploaded_pdf is not None:
        try:
            extracted_text = extract_pdf_text(uploaded_pdf)
            st.session_state["original_resume_text"] = extracted_text
            if extracted_text:
                with right:
                    st.subheader("Extracted resume text")
                    st.text_area("Review extracted text", extracted_text, height=300, key="extracted_preview")
            else:
                st.warning("No selectable text was found. If this is a scanned PDF, run OCR first.")
        except ValueError as exc:
            st.error(str(exc))

    if analyze_clicked:
        original = st.session_state.get("original_resume_text", "")
        if not original:
            st.error("Upload a readable PDF resume first.")
        elif not job_description.strip():
            st.error("Paste the job description first.")
        else:
            system_prompt = """
You are a careful resume and ATS analyst. Be evidence-based and concise.
Never invent skills, work experience, degrees, certifications, dates, achievements,
metrics, employers, or project outcomes. Distinguish supported, partially supported,
and missing requirements. Return ONLY valid JSON with this schema:
{
  "overall_match": 0,
  "summary": "short summary",
  "category_scores": [{"category":"Skills","score":0,"notes":"..."}],
  "matched_skills": ["..."],
  "partial_skills": [{"skill":"...","evidence":"...","gap":"..."}],
  "missing_skills": ["..."],
  "requirements_evidence": [{"requirement":"...","status":"Supported/Partial/Missing","evidence":"..."}],
  "recommendations": ["..."],
  "keyword_gaps": ["..."],
  "quality_checks": [{"check":"...","status":"Pass/Review","details":"..."}]
}
Scores must be estimates, not claims about a specific ATS vendor.
"""
            user_prompt = (
                "Analyze this resume against the job description.\n\n"
                "RESUME:\n" + original +
                "\n\nJOB DESCRIPTION:\n" + job_description +
                "\n\nADDITIONAL VERIFIED DETAILS:\n" + (verified_info or "None provided") +
                "\n\nApplicant context: " + ("Student/fresher" if fresher_mode else "General applicant")
            )
            with st.spinner("Analyzing resume..."):
                try:
                    raw = call_ai(system_prompt, user_prompt)
                    st.session_state["analysis_result"] = parse_json_response(raw)
                    st.session_state["last_job_description"] = job_description
                    st.session_state["last_verified_info"] = verified_info
                    st.success("Analysis complete.")
                except Exception as exc:
                    st.error(f"Could not complete the request: {exc}")

    analysis = st.session_state.get("analysis_result")
    if analysis:
        st.divider()
        st.header("Analysis Results")
        score = analysis.get("overall_match", "—")
        c1, c2 = st.columns([1, 3])
        with c1:
            st.metric("Estimated match", f"{score}%" if isinstance(score, (int, float)) else str(score))
        with c2:
            st.write(analysis.get("summary", ""))

        tab1, tab2, tab3, tab4 = st.tabs(
            ["Skills & Requirements", "Recommendations", "Tailored Resume", "Quality Checks"]
        )
        with tab1:
            st.subheader("Matched skills")
            display_list(analysis.get("matched_skills"))
            st.subheader("Partially matched skills")
            display_list(analysis.get("partial_skills"))
            st.subheader("Missing skills / keywords")
            display_list(analysis.get("missing_skills"))
            display_list(analysis.get("keyword_gaps"))
            evidence = safe_list(analysis.get("requirements_evidence"))
            if evidence:
                st.subheader("Requirement evidence")
                st.dataframe(evidence, use_container_width=True)
            category_scores = safe_list(analysis.get("category_scores"))
            if category_scores:
                st.subheader("Category breakdown")
                st.dataframe(category_scores, use_container_width=True)

        with tab2:
            st.subheader("Actionable recommendations")
            display_list(analysis.get("recommendations"))

        with tab3:
            st.subheader("Generate a tailored resume")
            st.caption("The draft uses only resume content and additional details you marked as verified.")
            if st.button("Generate tailored resume", type="primary"):
                original = st.session_state.get("original_resume_text", "")
                jd = st.session_state.get("last_job_description", "")
                verified = st.session_state.get("last_verified_info", "")
                resume_system = """
You are a professional resume editor. Create a clear, ATS-readable resume in Markdown.
Do not fabricate or infer any credentials, skills, experience, dates, metrics, awards,
links, or achievements. If information is absent, omit it or mark [ADD VERIFIED DETAIL].
Tailor wording and ordering to the job description while preserving truth.
Use concise bullets, standard section headings, and no tables or graphics.
"""
                resume_user = (
                    f"Create a tailored resume using the {template} style.\n"
                    f"Candidate name if provided: {candidate_name or 'Use name only if present in source'}\n"
                    f"Applicant context: {'Student/fresher' if fresher_mode else 'General applicant'}\n\n"
                    f"CURRENT RESUME:\n{original}\n\nJOB DESCRIPTION:\n{jd}\n\n"
                    f"ADDITIONAL VERIFIED DETAILS:\n{verified or 'None'}"
                )
                with st.spinner("Drafting tailored resume..."):
                    try:
                        tailored = call_ai(resume_system, resume_user)
                        st.session_state["tailored_resume"] = tailored
                    except Exception as exc:
                        st.error(f"Could not generate tailored resume: {exc}")

            tailored = st.session_state.get("tailored_resume", "")
            if tailored:
                edited = st.text_area(
                    "Edit your tailored resume before exporting",
                    value=tailored,
                    height=500,
                    key="tailored_resume_editor",
                )
                st.session_state["tailored_resume_final"] = edited
                st.download_button(
                    "Download as Markdown",
                    data=edited.encode("utf-8"),
                    file_name="tailored_resume.md",
                    mime="text/markdown",
                )
                if REPORTLAB_AVAILABLE:
                    try:
                        st.download_button(
                            "Download as PDF",
                            data=markdown_to_pdf(edited),
                            file_name="tailored_resume.pdf",
                            mime="application/pdf",
                        )
                    except Exception as exc:
                        st.warning(f"PDF export unavailable: {exc}")
                if DOCX_AVAILABLE:
                    try:
                        st.download_button(
                            "Download as DOCX",
                            data=markdown_to_docx(edited),
                            file_name="tailored_resume.docx",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        )
                    except Exception as exc:
                        st.warning(f"DOCX export unavailable: {exc}")

                save_name = st.text_input("Name this saved resume", value="Tailored Resume")
                if st.button("Save resume to local history"):
                    save_resume(save_name.strip() or "Tailored Resume", edited)
                    st.success("Saved locally in this app's SQLite database.")

                original = st.session_state.get("original_resume_text", "")
                with st.expander("Before / after text comparison"):
                    st.markdown("**Original extracted resume**")
                    st.text_area("Original", original, height=220, key="original_compare")
                    st.markdown("**Tailored draft**")
                    st.text_area("Tailored", edited, height=220, key="tailored_compare")

        with tab4:
            checks = safe_list(analysis.get("quality_checks"))
            if checks:
                st.dataframe(checks, use_container_width=True)
            else:
                st.info("No structured quality checks were returned.")


        st.divider()
        st.markdown("# 🚀 WHAT YOU NEED TO LEARN NEXT")
        st.markdown(
            "## **PERSONALIZED SKILL GAPS, LEARNING PRIORITIES & RESOURCES**"
        )
        st.write(
            "Use this section to turn missing or partial job requirements into a practical "
            "learning plan. Recommendations are generated from your resume analysis and "
            "the job description; they are learning suggestions, not claims that you already "
            "have these skills."
        )

        gap_items = []
        for item in safe_list(analysis.get("missing_skills")):
            gap_items.append(str(item.get("skill") or item.get("name") or item) if isinstance(item, dict) else str(item))
        for item in safe_list(analysis.get("partial_skills")):
            skill = item.get("skill") or item.get("name") or str(item) if isinstance(item, dict) else str(item)
            if skill and skill not in gap_items:
                gap_items.append(skill)

        # Presentation-only animation and controls; existing functionality remains intact.
        st.markdown("""
        <style>
        @keyframes riseIn { from {opacity:0; transform:translateY(14px)} to {opacity:1; transform:translateY(0)} }
        @keyframes softGlow { 0%,100% {box-shadow:0 0 0 rgba(80,150,255,0)} 50% {box-shadow:0 0 22px rgba(80,150,255,.16)} }
        .learn-hero {padding:1.25rem 1.4rem;border-radius:18px;margin:.5rem 0 1rem;border:1px solid rgba(100,160,255,.35);background:linear-gradient(120deg,rgba(30,90,180,.20),rgba(110,70,200,.12),rgba(20,150,170,.12));animation:riseIn .7s ease-out both,softGlow 4s ease-in-out infinite}
        .learn-card {border:1px solid rgba(130,160,200,.30);border-radius:16px;padding:1rem 1.1rem;margin:.65rem 0;background:linear-gradient(145deg,rgba(100,140,200,.10),rgba(140,100,200,.07));animation:riseIn .55s ease-out both;transition:transform .2s ease,border-color .2s ease,box-shadow .2s ease}
        .learn-card:hover {transform:translateY(-3px);border-color:rgba(100,170,255,.75);box-shadow:0 10px 28px rgba(40,90,170,.14)}
        .learn-kicker {font-size:.78rem;letter-spacing:.12em;text-transform:uppercase;opacity:.75;font-weight:700}
        @media (prefers-reduced-motion: reduce) {.learn-hero,.learn-card {animation:none;transition:none}}
        </style>
        <div class="learn-hero"><div class="learn-kicker">Personalized learning sprint</div>
        <h2 style="margin:.35rem 0">TURN SKILL GAPS INTO A DAILY CRASH COURSE</h2>
        <p style="margin:0;opacity:.88">Daily topics • hands-on tasks • free learning sources • progress checkpoints</p></div>
        """, unsafe_allow_html=True)

        crash_days = st.select_slider(
            "Crash course duration (days)", options=[3, 5, 7, 10, 14, 21, 30],
            value=7, key="crash_course_days"
        )
        daily_hours = st.select_slider(
            "Study time available per day",
            options=["30 minutes", "1 hour", "2 hours", "3+ hours"],
            value="1 hour", key="crash_course_daily_time"
        )

        if gap_items:
            st.markdown("### **YOUR IDENTIFIED LEARNING GAPS**")
            for gap in gap_items:
                st.markdown(f"- **{gap}**")
        else:
            st.info("The analysis did not identify missing or partially matched skills. You can still request a learning plan using the job description.")

        if st.button("🚀 BUILD MY DAY-BY-DAY CRASH COURSE", type="primary", key="generate_crash_course"):
            plan_resume = st.session_state.get("original_resume_text", "")
            plan_jd = st.session_state.get("last_job_description", "")
            plan_gaps = "\\n".join(f"- {x}" for x in gap_items) or "Identify cautious priorities from the resume and job description."
            crash_system = """
You are a practical career-learning coach. Create exactly the requested number of days.
Base the plan on the supplied resume, job description, and gaps. Never imply the learner
already knows a topic without evidence. For EVERY day provide: title, outcomes, ordered
topics, time estimate within the daily budget, hands-on practice, self-check/deliverable,
and recommended free source types. Prioritize essentials and label optional topics.
Finish with a capstone and readiness checklist. Do not invent exact course titles or
unverified direct video URLs; the app supplies topic-specific search links.
Use readable Markdown with clear day headings and concise bullets.
"""
            crash_user = (
                f"Make exactly {crash_days} days. Daily budget: {daily_hours}.\\n"
                f"Resume:\\n{plan_resume}\\n\\nJob description:\\n{plan_jd}\\n\\n"
                f"Missing/partial skills:\\n{plan_gaps}\\n\\n"
                f"Additional verified details:\\n{st.session_state.get('last_verified_info', 'None provided')}"
            )
            with st.spinner("Designing your daily crash course..."):
                try:
                    st.session_state["daily_crash_course"] = call_ai(crash_system, crash_user)
                    st.session_state["crash_course_gap_list"] = gap_items
                    st.session_state["crash_course_duration"] = crash_days
                except Exception as exc:
                    st.error(f"Could not generate crash course: {exc}")

        if st.session_state.get("daily_crash_course"):
            st.markdown("## **YOUR DAY-BY-DAY CRASH COURSE**")
            st.markdown(st.session_state["daily_crash_course"])
            st.markdown("## **FREE LEARNING HUB — CHOOSE YOUR SOURCE**")
            st.caption("Topic-specific search links help you choose free materials at your level. Check access/pricing before starting.")
            from urllib.parse import quote_plus
            course_topics = st.session_state.get("crash_course_gap_list", gap_items)
            if not course_topics:
                course_topics = ["the key skills for this job description"]
            for idx, topic in enumerate(course_topics):
                topic_text = str(topic)
                q = quote_plus(topic_text + " beginner tutorial")
                doc_q = quote_plus(topic_text + " official documentation")
                st.markdown(
                    f'<div class="learn-card"><div class="learn-kicker">Learning focus {idx+1}</div>'
                    f'<h3 style="margin:.25rem 0 .6rem">{topic_text}</h3>'
                    f'<p><a href="https://www.youtube.com/results?search_query={q}" target="_blank">▶ YouTube lessons</a>'
                    f' &nbsp;·&nbsp; <a href="https://www.freecodecamp.org/news/search/?query={q}" target="_blank">freeCodeCamp</a>'
                    f' &nbsp;·&nbsp; <a href="https://www.google.com/search?q={doc_q}" target="_blank">Official docs</a>'
                    f' &nbsp;·&nbsp; <a href="https://www.coursera.org/search?query={q}" target="_blank">Coursera (check free options)</a>'
                    f' &nbsp;·&nbsp; <a href="https://github.com/search?q={q}&type=repositories" target="_blank">GitHub practice</a></p></div>',
                    unsafe_allow_html=True
                )
            st.markdown("### **DAILY PROGRESS CHECKLIST**")
            for day_num in range(1, int(st.session_state.get("crash_course_duration", crash_days)) + 1):
                st.checkbox(f"Day {day_num}: complete lessons, practice, and self-check", key=f"crash_course_done_{day_num}")

        if st.button("Generate my learning recommendations", type="primary", key="generate_learning_recs"):
            original = st.session_state.get("original_resume_text", "")
            jd = st.session_state.get("last_job_description", "")
            gaps_text = "\n".join(f"- {x}" for x in gap_items) or "No explicit gaps listed; infer learning priorities cautiously from the job description and resume."
            learning_system = """
You are a practical technical learning mentor. Analyze the candidate's resume, job description,
and detected skill gaps. Recommend only skills that are genuinely missing, partially evidenced,
or important prerequisites. Do not say the candidate knows something unless the resume supports it.
Return readable Markdown with:
1. A bold top-priority summary.
2. A table or clear numbered list: skill/topic to learn, why it matters for this role,
   beginner-to-intermediate topics in order, a small hands-on practice task, and a completion check.
3. A realistic sequence (start here, next, then).
4. Clearly distinguish required gaps from optional/advanced topics.
Avoid pretending to have checked the candidate's actual knowledge beyond the supplied resume.
Do not invent specific course titles or URLs; the app adds resource search links separately.
"""
            learning_user = (
                "RESUME:\n" + original +
                "\n\nJOB DESCRIPTION:\n" + jd +
                "\n\nIDENTIFIED MISSING/PARTIAL SKILLS:\n" + gaps_text +
                "\n\nADDITIONAL VERIFIED DETAILS:\n" +
                st.session_state.get("last_verified_info", "None provided")
            )
            with st.spinner("Building your personalized learning plan..."):
                try:
                    st.session_state["learning_recommendations"] = call_ai(learning_system, learning_user)
                    st.session_state["learning_gap_list"] = gap_items
                except Exception as exc:
                    st.error(f"Could not generate learning recommendations: {exc}")

        learning_recommendations = st.session_state.get("learning_recommendations", "")
        if learning_recommendations:
            st.markdown(learning_recommendations)
            st.markdown("## **WHERE TO LEARN — MULTIPLE SOURCES**")
            st.caption("These are topic-specific search links, so you can choose a tutorial that matches your level. Check the course date, prerequisites, and whether it is free before starting.")
            from urllib.parse import quote_plus

            resource_gaps = st.session_state.get("learning_gap_list", gap_items)
            for gap in resource_gaps:
                q = quote_plus(gap + " tutorial for beginners")
                st.markdown(f"### **{gap}**")
                st.markdown(
                    f"- [YouTube tutorials](https://www.youtube.com/results?search_query={q})\n"
                    f"- [freeCodeCamp learning resources](https://www.freecodecamp.org/news/search/?query={q})\n"
                    f"- [Official documentation / guides (Google search)](https://www.google.com/search?q={quote_plus(gap + ' official documentation tutorial')})\n"
                    f"- [Structured courses (Coursera search)](https://www.coursera.org/search?query={q})\n"
                    f"- [Hands-on projects and examples (GitHub search)](https://github.com/search?q={q}&type=repositories)"
                )
            st.markdown("## **YOUR NEXT ACTIONS**")
            st.checkbox("Choose one priority skill and bookmark one learning resource.", key="learning_action_1")
            st.checkbox("Complete a small hands-on exercise without copying the full solution.", key="learning_action_2")
            st.checkbox("Write down what you learned and update your resume only after you can demonstrate it.", key="learning_action_3")


# ----------------------------- Interview Prep -------------------------------
elif page == "Interview Preparation":
    st.header("Interview Preparation")
    role = st.text_input("Target role", placeholder="e.g., Python Developer Intern")
    jd = st.text_area("Job description (optional)", height=160)
    resume = st.text_area("Relevant resume / project details (optional)", height=160)
    count = st.slider("Number of questions", 3, 15, 8)
    if st.button("Generate interview practice", type="primary"):
        if not role.strip():
            st.warning("Enter a target role.")
        else:
            system = (
                "You are an interview coach. Generate role-relevant practice questions "
                "and concise guidance on what a strong answer should cover. Do not claim "
                "the user has experience not included in their details."
            )
            user = (
                f"Role: {role}\nJob description: {jd or 'Not provided'}\n"
                f"Candidate details: {resume or 'Not provided'}\n"
                f"Create {count} questions, grouped into technical, project, and behavioral "
                "where appropriate. Include brief answer guidance."
            )
            with st.spinner("Preparing questions..."):
                try:
                    st.session_state["interview_pack"] = call_ai(system, user)
                except Exception as exc:
                    st.error(f"Could not generate questions: {exc}")
    if st.session_state.get("interview_pack"):
        st.markdown(st.session_state["interview_pack"])
        st.divider()
        st.subheader("Practice an answer")
        question = st.text_area("Paste or write one interview question")
        answer = st.text_area("Your answer", height=180)
        if st.button("Get feedback"):
            if not question.strip() or not answer.strip():
                st.warning("Provide both a question and your answer.")
            else:
                with st.spinner("Reviewing your answer..."):
                    try:
                        feedback = call_ai(
                            "Give constructive interview-answer feedback. Identify strengths, "
                            "gaps, and a suggested structure. Do not write false personal claims.",
                            f"Question: {question}\nCandidate answer: {answer}"
                        )
                        st.markdown(feedback)
                    except Exception as exc:
                        st.error(f"Could not review answer: {exc}")


# ----------------------------- Skill Roadmap -------------------------------
elif page == "Skill Roadmap":
    st.header("Personalized Skill Roadmap")
    target = st.text_input("Target role", placeholder="e.g., Backend Developer Intern")
    current = st.text_area("Current skills and experience", height=140)
    available_time = st.selectbox("Weekly time available", ["2–4 hours", "5–7 hours", "8–10 hours", "10+ hours"])
    if st.button("Build my roadmap", type="primary"):
        if not target.strip():
            st.warning("Enter a target role.")
        else:
            with st.spinner("Building roadmap..."):
                try:
                    roadmap = call_ai(
                        "You are a practical learning mentor. Make a realistic beginner-friendly "
                        "roadmap with milestones, practice tasks, and measurable completion criteria. "
                        "Do not assume prior knowledge not stated.",
                        f"Target role: {target}\nCurrent background: {current or 'Not provided'}\n"
                        f"Time available: {available_time}\nCreate a staged 6-8 week plan."
                    )
                    st.session_state["skill_roadmap"] = roadmap
                except Exception as exc:
                    st.error(f"Could not create roadmap: {exc}")
    if st.session_state.get("skill_roadmap"):
        st.markdown(st.session_state["skill_roadmap"])
        st.caption("Use these checkboxes to track milestones locally during this session.")
        for i, line in enumerate(
            [line.strip("- *") for line in st.session_state["skill_roadmap"].splitlines()
             if line.strip() and (line.lstrip().startswith(("-", "*")) or re.match(r"^\d+[.)]", line.lstrip()))][:30]
        ):
            st.checkbox(line, key=f"roadmap_check_{i}")


# ----------------------------- Application Tracker --------------------------
elif page == "Application Tracker":
    st.header("Job Application Tracker")
    with st.form("add_application_form", clear_on_submit=True):
        company = st.text_input("Company")
        role = st.text_input("Role")
        status = st.selectbox("Status", ["Planning", "Applied", "Assessment", "Interview", "Offer", "Rejected", "Withdrawn"])
        date_applied = st.date_input("Date applied", value=datetime.now().date())
        notes = st.text_area("Notes")
        submitted = st.form_submit_button("Add application")
        if submitted:
            if company.strip() and role.strip():
                add_application(company.strip(), role.strip(), status, str(date_applied), notes.strip())
                st.success("Application added.")
            else:
                st.error("Company and role are required.")

    applications = get_applications()
    if applications:
        st.subheader("Your applications")
        for app in applications:
            app_id, company, role, status, date_applied, notes = app
            with st.expander(f"{company} — {role} | {status}"):
                st.write(f"Date applied: {date_applied or '—'}")
                new_status = st.selectbox(
                    "Update status", ["Planning", "Applied", "Assessment", "Interview", "Offer", "Rejected", "Withdrawn"],
                    index=["Planning", "Applied", "Assessment", "Interview", "Offer", "Rejected", "Withdrawn"].index(status)
                    if status in ["Planning", "Applied", "Assessment", "Interview", "Offer", "Rejected", "Withdrawn"] else 0,
                    key=f"app_status_{app_id}"
                )
                new_notes = st.text_area("Update notes", value=notes or "", key=f"app_notes_{app_id}")
                col_a, col_b = st.columns(2)
                with col_a:
                    if st.button("Save changes", key=f"save_app_{app_id}"):
                        update_application(app_id, new_status, new_notes)
                        st.success("Updated. Refresh if needed.")
                with col_b:
                    if st.button("Delete application", key=f"delete_app_{app_id}"):
                        delete_application(app_id)
                        st.rerun()
    else:
        st.info("No applications added yet.")


# ----------------------------- Saved Resumes -------------------------------
elif page == "Saved Resumes":
    st.header("Saved Resume History")
    resumes = get_saved_resumes()
    if not resumes:
        st.info("No saved resumes yet. Generate and save one from the Resume Analyzer page.")
    else:
        for resume_id, name, created_at in resumes:
            with st.expander(f"{name} — {created_at}"):
                text = get_resume_text(resume_id)
                st.text_area("Saved resume", text, height=260, key=f"saved_text_{resume_id}")
                col1, col2, col3 = st.columns(3)
                with col1:
                    st.download_button(
                        "Download Markdown",
                        data=text.encode("utf-8"),
                        file_name=f"resume_{resume_id}.md",
                        mime="text/markdown",
                        key=f"download_md_{resume_id}",
                    )
                with col2:
                    if REPORTLAB_AVAILABLE:
                        st.download_button(
                            "Download PDF",
                            data=markdown_to_pdf(text),
                            file_name=f"resume_{resume_id}.pdf",
                            mime="application/pdf",
                            key=f"download_pdf_{resume_id}",
                        )
                with col3:
                    if DOCX_AVAILABLE:
                        st.download_button(
                            "Download DOCX",
                            data=markdown_to_docx(text),
                            file_name=f"resume_{resume_id}.docx",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            key=f"download_docx_{resume_id}",
                        )
                if st.button("Delete saved resume", key=f"delete_resume_{resume_id}"):
                    delete_resume(resume_id)
                    st.rerun()

st.sidebar.divider()
st.sidebar.caption("Privacy note: resumes and tracker entries are stored in a local SQLite file beside app.py. Avoid uploading sensitive information.")
